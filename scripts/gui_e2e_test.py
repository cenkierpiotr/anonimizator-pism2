"""Test e2e GUI pod Xvfb: dodanie realnego PDF-a z zawijanym tekstem,
uruchomienie anonimizacji, przejscie ekranu weryfikacji, zapis wyniku -
klikajac PRAWDZIWE widgety (xdotool), nie wywolujac metod bezposrednio.
Jedyne monkeypatche to natywne dialogi systemowe (askopenfilenames/
asksaveasfilename/messagebox) - to okna OS poza drzewem widgetow aplikacji,
ktorych nie da sie ani wyrenderowac w zrzucie apki, ani sensownie
zautomatyzowac bez zaleznosci od lokalizacji/motywu systemu."""

from __future__ import annotations

import subprocess
import sys
import time
from pathlib import Path
from tkinter import filedialog, messagebox

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fpdf import FPDF
import docx

SCREEN_DIR = Path("/tmp/gui_e2e_screens")
SCREEN_DIR.mkdir(exist_ok=True)

INPUT_PDF = Path("/tmp/gui_e2e_input.pdf")
OUTPUT_DOCX = Path("/tmp/gui_e2e_output.docx")


def make_test_pdf() -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    p1 = (
        "W dniu 15 stycznia 2026 roku pozwany Jan Kowalski zamieszkaly w Warszawie "
        "przy ulicy Marszalkowskiej 10 zawarl umowe najmu lokalu z powodem Anna Nowak "
        "na okres dwunastu miesiecy za czynsz w wysokosci 3000 zlotych miesiecznie. "
        "Numer telefonu powoda to 601 234 567, a adres e-mail anna.nowak@example.com."
    )
    p2 = (
        "Powod wnosi o zaplate zaleglego czynszu w kwocie 9000 zlotych wraz z odsetkami "
        "ustawowymi za opoznienie liczonymi od dnia wniesienia pozwu do dnia zaplaty."
    )
    pdf.multi_cell(0, 8, p1)
    pdf.ln(8)
    pdf.multi_cell(0, 8, p2)
    pdf.output(str(INPUT_PDF))


def xdotool(*args: str) -> str:
    return subprocess.run(["xdotool", *args], capture_output=True, text=True).stdout.strip()


def click_widget(widget) -> None:
    widget.update_idletasks()
    x = widget.winfo_rootx() + widget.winfo_width() // 2
    y = widget.winfo_rooty() + widget.winfo_height() // 2
    xdotool("mousemove", "--sync", str(x), str(y))
    xdotool("click", "1")


def screenshot(name: str) -> None:
    path = SCREEN_DIR / f"{name}.xwd"
    subprocess.run(["bash", "-c", f"xwd -root -display $DISPLAY -out '{path}'"], check=False)
    print(f"screenshot: {path} (exists={path.exists()})")


def main() -> int:
    make_test_pdf()
    OUTPUT_DOCX.unlink(missing_ok=True)

    filedialog.askopenfilenames = lambda **kwargs: (str(INPUT_PDF),)
    filedialog.asksaveasfilename = lambda **kwargs: str(OUTPUT_DOCX)
    messagebox.showinfo = lambda *a, **k: print(f"[messagebox.showinfo] {a}")
    messagebox.showerror = lambda *a, **k: print(f"[messagebox.showerror] {a}")

    from app.gui.main_window import AnonymizerApp, FileStatus

    app = AnonymizerApp()
    app.update_idletasks()
    app.update()
    screenshot("1_initial")

    assert hasattr(app, "cancel_button"), "Brak przycisku Anuluj w interfejsie"
    toolbar_buttons = {}
    for child in app.winfo_children()[0].winfo_children():
        text = child.cget("text") if hasattr(child, "cget") else None
        if text:
            toolbar_buttons[text] = child
    print("Przyciski toolbara znalezione:", list(toolbar_buttons.keys()))
    expected_buttons = {
        "Dodaj pliki...",
        "Anonimizuj zaznaczone",
        "Konwertuj do PDF",
        "Zainstaluj obsługę .doc",
        "Anuluj",
        "Usuń zaznaczone z listy",
    }
    missing = expected_buttons - set(toolbar_buttons.keys())
    assert not missing, f"Brakujące przyciski w toolbarze: {missing}"

    click_widget(toolbar_buttons["Dodaj pliki..."])
    app.update()
    time.sleep(0.2)
    app.update()
    assert len(app.items) == 1, f"Oczekiwano 1 pliku w kolejce, jest {len(app.items)}"
    print("Plik dodany do kolejki:", list(app.items.keys()))
    screenshot("2_after_add")

    click_widget(toolbar_buttons["Anonimizuj zaznaczone"])
    app.update()

    deadline = time.time() + 60
    key = list(app.items.keys())[0]
    while time.time() < deadline:
        app.update()
        time.sleep(0.1)
        if app.items[key].status == FileStatus.DO_WERYFIKACJI:
            break
    else:
        raise AssertionError(f"Timeout - status pozostal: {app.items[key].status}")

    print("Status po przetworzeniu:", app.items[key].status)
    screenshot("3_after_processing")

    staged = app.items[key].staged
    assert staged is not None
    print("Ostrzezenia:", staged.warnings)
    print("Liczba wykrytych encji:", staged.entity_count)
    assert staged.entity_count >= 4, f"Za malo wykrytych encji: {staged.entity_count}"

    row_widgets = app.row_widgets[key]
    click_widget(row_widgets["review_button"])
    app.update()
    time.sleep(0.2)
    app.update()
    screenshot("4_review_window")

    review_windows = [w for w in app.winfo_children() if "toplevel" in str(w).lower() or w.winfo_class() == "Toplevel"]
    print("Toplevel windows:", review_windows)

    approve_button = None
    for w in app.winfo_children():
        if w.winfo_class() != "Toplevel":
            continue
        for child in w.winfo_children():
            if hasattr(child, "winfo_children"):
                for grandchild in child.winfo_children():
                    if hasattr(grandchild, "cget"):
                        try:
                            if grandchild.cget("text") == "Zatwierdź i zapisz...":
                                approve_button = grandchild
                        except Exception:
                            pass
    assert approve_button is not None, "Nie znaleziono przycisku 'Zatwierdź i zapisz...' w oknie recenzji"

    click_widget(approve_button)
    app.update()
    time.sleep(0.3)
    app.update()
    screenshot("5_after_approve")

    assert app.items[key].status == FileStatus.ZAPISANY, f"Status koncowy: {app.items[key].status}"
    assert OUTPUT_DOCX.exists(), "Plik wyjsciowy nie zostal zapisany"

    document = docx.Document(str(OUTPUT_DOCX))
    print(f"--- Wynikowy dokument: {len(document.paragraphs)} akapitow ---")
    for i, p in enumerate(document.paragraphs):
        print(f"[{i}] {p.text!r}")

    assert len(document.paragraphs) == 2, f"Oczekiwano 2 akapitow, jest {len(document.paragraphs)}"
    full_text = "\n".join(p.text for p in document.paragraphs)
    assert "Jan Kowalski" not in full_text
    assert "601 234 567" not in full_text
    assert "anna.nowak@example.com" not in full_text
    assert "[Osoba 1]" in full_text
    assert "[adres 1]" in full_text

    app.destroy()
    print("\nOK: pelny przeplyw GUI dziala poprawnie, tekst nie jest urywany.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
