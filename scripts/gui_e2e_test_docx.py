"""Test e2e GUI dla wejscia .docx (analogiczny do gui_e2e_test.py, ktory
pokrywa tylko PDF): dodanie realnego .docx z PII wielowatkowo przez ta sama
sciezke kodu co PDF (docx_reader -> detect_all -> docx_writer), zeby
potwierdzic ze caly przeplyw dodaj->anonimizuj->recenzja->zapisz dziala
rowniez dla najczestszego w praktyce formatu wejsciowego."""

from __future__ import annotations

import subprocess
import sys
import time
from pathlib import Path
from tkinter import filedialog, messagebox

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import docx

SCREEN_DIR = Path("/tmp/gui_e2e_screens_docx")
SCREEN_DIR.mkdir(exist_ok=True)

INPUT_DOCX = Path("/tmp/gui_e2e_input.docx")
OUTPUT_DOCX = Path("/tmp/gui_e2e_output_docx.docx")


def make_test_docx() -> None:
    document = docx.Document()
    document.add_paragraph(
        "Powod Jan Kowalski, PESEL 44051401359, zamieszkaly w Warszawie przy "
        "ulicy Marszalkowskiej 10, tel. 601-234-567, e-mail jan.kowalski@example.com, "
        "wnosi pozew przeciwko Annie Nowak o zaplate kwoty 9000 zlotych."
    )
    document.add_paragraph(
        "Pozwana Anna Nowak zamieszkala jest przy ulicy Marszalkowskiej 10 w Warszawie."
    )
    document.save(INPUT_DOCX)


def xdotool(*args: str) -> str:
    return subprocess.run(["xdotool", *args], capture_output=True, text=True).stdout.strip()


def click_widget(widget) -> None:
    widget.update_idletasks()
    x = widget.winfo_rootx() + widget.winfo_width() // 2
    y = widget.winfo_rooty() + widget.winfo_height() // 2
    xdotool("mousemove", "--sync", str(x), str(y))
    xdotool("click", "1")


def screenshot(name: str) -> None:
    path = SCREEN_DIR / f"{name}.xwd"
    subprocess.run(["bash", "-c", f"xwd -root -display $DISPLAY -out '{path}'"], check=False)
    print(f"screenshot: {path} (exists={path.exists()})")


def main() -> int:
    make_test_docx()
    OUTPUT_DOCX.unlink(missing_ok=True)

    filedialog.askopenfilenames = lambda **kwargs: (str(INPUT_DOCX),)
    filedialog.asksaveasfilename = lambda **kwargs: str(OUTPUT_DOCX)
    messagebox.showinfo = lambda *a, **k: print(f"[messagebox.showinfo] {a}")
    messagebox.showerror = lambda *a, **k: print(f"[messagebox.showerror] {a}")

    from app.gui.main_window import AnonymizerApp, FileStatus

    app = AnonymizerApp()
    app.update_idletasks()
    app.update()
    screenshot("1_initial")

    toolbar_buttons = {}
    for child in app.winfo_children()[0].winfo_children():
        text = child.cget("text") if hasattr(child, "cget") else None
        if text:
            toolbar_buttons[text] = child

    click_widget(toolbar_buttons["Dodaj pliki..."])
    app.update()
    time.sleep(0.2)
    app.update()
    assert len(app.items) == 1, f"Oczekiwano 1 pliku w kolejce, jest {len(app.items)}"
    print("Plik dodany do kolejki:", list(app.items.keys()))
    screenshot("2_after_add")

    click_widget(toolbar_buttons["Anonimizuj zaznaczone"])
    app.update()

    deadline = time.time() + 60
    key = list(app.items.keys())[0]
    while time.time() < deadline:
        app.update()
        time.sleep(0.1)
        if app.items[key].status == FileStatus.DO_WERYFIKACJI:
            break
    else:
        raise AssertionError(f"Timeout - status pozostal: {app.items[key].status}")

    print("Status po przetworzeniu:", app.items[key].status)
    screenshot("3_after_processing")

    staged = app.items[key].staged
    assert staged is not None
    print("Ostrzezenia:", staged.warnings)
    print("Liczba wykrytych encji:", staged.entity_count)
    assert staged.entity_count >= 5, f"Za malo wykrytych encji: {staged.entity_count}"

    row_widgets = app.row_widgets[key]
    click_widget(row_widgets["review_button"])
    app.update()
    time.sleep(0.2)
    app.update()
    screenshot("4_review_window")

    approve_button = None
    for w in app.winfo_children():
        if w.winfo_class() != "Toplevel":
            continue
        for child in w.winfo_children():
            if hasattr(child, "winfo_children"):
                for grandchild in child.winfo_children():
                    if hasattr(grandchild, "cget"):
                        try:
                            if grandchild.cget("text") == "Zatwierdź i zapisz...":
                                approve_button = grandchild
                        except Exception:
                            pass
    assert approve_button is not None, "Nie znaleziono przycisku 'Zatwierdź i zapisz...' w oknie recenzji"

    click_widget(approve_button)
    app.update()
    time.sleep(0.3)
    app.update()
    screenshot("5_after_approve")

    assert app.items[key].status == FileStatus.ZAPISANY, f"Status koncowy: {app.items[key].status}"
    assert OUTPUT_DOCX.exists(), "Plik wyjsciowy nie zostal zapisany"

    document = docx.Document(str(OUTPUT_DOCX))
    print(f"--- Wynikowy dokument: {len(document.paragraphs)} akapitow ---")
    for i, p in enumerate(document.paragraphs):
        print(f"[{i}] {p.text!r}")

    full_text = "\n".join(p.text for p in document.paragraphs)
    assert "Jan Kowalski" not in full_text
    assert "44051401359" not in full_text
    assert "601-234-567" not in full_text
    assert "jan.kowalski@example.com" not in full_text
    assert "9000" in full_text, "Kwota nie powinna byc anonimizowana (polityka: zostaw)"
    assert "[Osoba 1]" in full_text
    assert "[adres 1]" in full_text
    assert "[PESEL 1]" in full_text

    app.destroy()
    print("\nOK: pelny przeplyw GUI dla .docx dziala poprawnie.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
