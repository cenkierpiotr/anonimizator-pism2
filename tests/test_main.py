import zipfile

import docx
import pytest

from app.main import AnonymizeOptions, anonymize_file, anonymize_files, discard_staged, process_to_staging
from app.pipeline.docx_writer import W_NS


@pytest.fixture
def sample_docx(tmp_path):
    path = tmp_path / "pozew.docx"
    document = docx.Document()
    document.add_paragraph(
        "Powód Jan Kowalski, PESEL 44051401359, tel. 601-234-567, "
        "wnosi pozew przeciwko Annie Nowak."
    )
    document.save(path)
    return path


@pytest.fixture
def sample_txt(tmp_path):
    path = tmp_path / "pismo.txt"
    path.write_text(
        "Wnioskodawca PESEL 44051401359 wnosi o rozpoznanie sprawy.\n\n"
        "Kontakt: 601-234-567.",
        encoding="utf-8",
    )
    return path


def test_anonymize_docx_end_to_end(sample_docx, tmp_path):
    output = tmp_path / "out" / "wynik.docx"
    result = anonymize_file(sample_docx, output)

    assert output.exists()
    assert result.entity_count > 0

    document = docx.Document(output)
    full_text = "\n".join(p.text for p in document.paragraphs)
    assert "44051401359" not in full_text
    assert "601-234-567" not in full_text
    assert "[PESEL 1]" in full_text
    assert "[numer telefonu 1]" in full_text


def test_anonymize_txt_end_to_end(sample_txt, tmp_path):
    output = tmp_path / "wynik.docx"
    result = anonymize_file(sample_txt, output)

    assert output.exists()
    document = docx.Document(output)
    full_text = "\n".join(p.text for p in document.paragraphs)
    assert "44051401359" not in full_text
    assert "601-234-567" not in full_text
    assert "[PESEL 1]" in full_text


def test_anonymize_files_queue_continues_after_single_failure(sample_docx, tmp_path):
    bad_input = tmp_path / "nieznany.xyz"
    bad_input.write_bytes(b"\x00\x01garbage")
    output_dir = tmp_path / "out"

    results = anonymize_files([sample_docx, bad_input], output_dir)

    assert len(results) == 2
    good_input, good_result, good_error = results[0]
    assert good_result is not None
    assert good_error is None
    assert good_result.output_path.exists()

    bad_path, bad_result, bad_error = results[1]
    assert bad_result is None
    assert bad_error is not None


def _add_document_protection(docx_path):
    settings_xml = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:settings xmlns:w="{W_NS}">'
        f'<w:documentProtection w:edit="readOnly" w:enforcement="1"/>'
        f"</w:settings>"
    ).encode("utf-8")

    tmp_path = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(
        tmp_path, "w", zipfile.ZIP_DEFLATED
    ) as zout:
        written_settings = False
        for item in zin.infolist():
            if item.filename == "word/settings.xml":
                zout.writestr(item, settings_xml)
                written_settings = True
            else:
                zout.writestr(item, zin.read(item.filename))
        if not written_settings:
            zout.writestr("word/settings.xml", settings_xml)
    tmp_path.replace(docx_path)


def test_document_protection_stripped_with_warning(sample_docx, tmp_path):
    _add_document_protection(sample_docx)
    output = tmp_path / "wynik.docx"

    result = anonymize_file(sample_docx, output)

    assert any("ochron" in w.lower() for w in result.warnings)
    with zipfile.ZipFile(output, "r") as zf:
        settings_xml = zf.read("word/settings.xml")
    assert b"documentProtection" not in settings_xml


def test_foreign_language_document_gets_warning(tmp_path):
    path = tmp_path / "letter.docx"
    document = docx.Document()
    document.add_paragraph(
        "This is a plain English letter with no sensitive data at all, "
        "just some ordinary words to fill out the paragraph nicely. "
        "The purpose of this text is simply to provide enough content "
        "for the language detection heuristic to reach a firm decision "
        "about which language this document was written in."
    )
    document.save(path)
    output = tmp_path / "wynik.docx"

    result = anonymize_file(path, output)

    assert any("obcojęzyczny" in w for w in result.warnings)


def test_polish_document_has_no_language_warning(sample_docx, tmp_path):
    output = tmp_path / "wynik.docx"

    result = anonymize_file(sample_docx, output)

    assert not any("obcojęzyczny" in w for w in result.warnings)


def test_staging_reports_potentially_missed_token(tmp_path):
    path = tmp_path / "pismo.docx"
    document = docx.Document()
    document.add_paragraph(
        "Spotkanie miało miejsce w Poniedziałek po południu w biurze."
    )
    document.save(path)

    staged = process_to_staging(path)
    try:
        assert any("Poniedziałek" in snippet for snippet in staged.potentially_missed)
    finally:
        discard_staged(staged)
